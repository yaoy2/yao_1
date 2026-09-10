import tempfile
import unittest
from pathlib import Path

from docx import Document

from utils.schedule_docx import WEEKDAYS, is_physical_education, parse_course_cell, parse_teacher_timetable, parse_weeks


def course_text(name="医学基础", code="123", weeks="1-16周", room="F3201", count=30, classes="医学影像25201"):
    return f"{name}[{code}] 3.0学分\n{weeks}[讲授] {room}\n选课人数: {count}\n{classes}"


def add_teacher(document, name):
    document.add_paragraph(f"2026-2027学年第1学期  {name}[A123456] 课表")
    table = document.add_table(rows=15, cols=6)
    for cell, text in zip(table.rows[0].cells, ["节次/星期", *WEEKDAYS]):
        cell.text = text
    for period in range(1, 15):
        table.cell(period, 0).text = f"第{period}节\n(08:20-09:00)"
    return table


class ScheduleDocxTest(unittest.TestCase):
    def test_disjoint_weeks_and_even_odd_weeks(self):
        self.assertEqual([2, 3, 5, 6, 8, 9], parse_weeks("2-3周,5-6周,8-9周"))
        self.assertEqual([6, 8, 10, 12], parse_weeks("6-12周(双)"))
        self.assertEqual([1, 3, 5], parse_weeks("1-6周（单）"))
        for invalid in ["16-1周", "周次待定", "0-16周"]:
            with self.assertRaises(ValueError):
                parse_weeks(invalid)

    def test_keeps_same_course_week_and_room_pairings(self):
        first = course_text(weeks="1-8周", room="F3102", classes="健康管理24201")
        second = course_text(weeks="9-16周", room="F3201", classes="健康管理25203")
        records = parse_course_cell(first + "\n" + second)
        self.assertEqual(2, len(records))
        self.assertEqual([("1-8周", "F3102", "健康管理24201"), ("9-16周", "F3201", "健康管理25203")], [
            (r["week_text"], r["classroom"], r["class_group"]) for r in records
        ])

    def test_combines_week_segments_only_for_same_teaching_type_and_room(self):
        text = "医学心理学[123] 2.0学分\n3-4周[讲授] F3118生理学实验室\n5-6周[讲授] F3118生理学实验室\n7-8周[实践] F3222\n选课人数: 74\n医学影像24205"
        records = parse_course_cell(text)
        self.assertEqual(2, len(records))
        self.assertEqual([3, 4, 5, 6], records[0]["weeks"])
        self.assertEqual("F3118生理学实验室", records[0]["classroom"])
        self.assertEqual([7, 8], records[1]["weeks"])
        self.assertEqual("实践", records[1]["teaching_type"])

    def test_retains_truncated_courses_without_guessing_missing_fields(self):
        complete = course_text()
        partial = "体育与健康（一）[1200500110] 1.0学..."
        records = parse_course_cell(complete + "\n" + partial)
        self.assertFalse(records[0]["source_truncated"])
        self.assertTrue(records[1]["source_truncated"])
        self.assertEqual("体育与健康（一）", records[1]["course"])
        self.assertEqual([], records[1]["weeks"])
        self.assertIsNone(records[1]["credits"])
        missing_class = parse_course_cell("流行病学[123] 2.0学分\n1-2周[讲授] F3202\n选课人数: 29...")[0]
        self.assertEqual("", missing_class["class_group"])
        self.assertTrue(missing_class["source_truncated"])

    def test_empty_classroom_and_zero_students_are_not_fabricated(self):
        record = parse_course_cell(course_text(room="", count=0, classes=""))[0]
        self.assertEqual("", record["classroom"])
        self.assertEqual("", record["class_group"])
        self.assertEqual(0, record["student_count"])
        self.assertFalse(record["source_truncated"])

    def test_rejects_unrecognized_cell_instead_of_silently_dropping_it(self):
        with self.assertRaises(ValueError):
            parse_course_cell("尚未支持的课程表格式")

    def test_import_preserves_merged_late_periods_and_excludes_all_pe_courses(self):
        document = Document()
        first = add_teacher(document, "张三")
        first.cell(6, 4).merge(first.cell(8, 4)).text = course_text(weeks="1-8周") + "\n" + course_text(weeks="9-16周", room="F3202")
        first.cell(11, 2).merge(first.cell(14, 2)).text = course_text(name="健康信息传播", code="456")
        mixed = add_teacher(document, "李四")
        mixed.cell(1, 1).merge(mixed.cell(2, 1)).text = course_text(name="体育（三）")
        mixed.cell(3, 1).merge(mixed.cell(4, 1)).text = course_text(name="大学生就业指导")
        sports_only = add_teacher(document, "王五")
        sports_only.cell(6, 1).merge(sports_only.cell(7, 1)).text = course_text(name="Physical Education and Health (I)")
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "测试课表.docx"
            document.save(path)
            records, metadata = parse_teacher_timetable(path, "2026-2027-1")
            with self.assertRaisesRegex(ValueError, "学期"):
                parse_teacher_timetable(path, "2025-2026-2")
        self.assertEqual(4, len(records))
        self.assertEqual(["张三", "李四"], metadata["teachers"])
        self.assertEqual(2, metadata["excluded_physical_education_records"])
        self.assertEqual(3, metadata["source_teacher_count"])
        self.assertEqual(13, metadata["source_occupied_period_cells"])
        late = next(r for r in records if r["course_code"] == "456")
        self.assertEqual(("星期二", 11, 14), (late["weekday"], late["start_period"], late["end_period"]))
        self.assertEqual("张三", late["teacher_label"])
        self.assertEqual(2, sum(r["weekday"] == "星期四" for r in records))

    def test_sports_filter_preserves_health_and_medical_courses(self):
        for name in ["体育（三）", "体育与健康（一）", "Physical Education（3）", "Physical Education and Health (I)"]:
            self.assertTrue(is_physical_education(name))
        for name in ["健康照护", "常见健康问题的识别与日常管理", "医学心理学"]:
            self.assertFalse(is_physical_education(name))


if __name__ == "__main__":
    unittest.main()
