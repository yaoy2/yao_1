from copy import deepcopy
from io import BytesIO
import base64
import json
from zipfile import ZipFile
from types import SimpleNamespace
from pathlib import Path
import ast
import hashlib
import hmac
import time

import pytest

from utils import title_review as review
from utils import title_review_sync as cloud
from utils.data_sync_validation import MANAGED_FILES, validate_backup, would_drop_records


SETTINGS = {"normal_scheme_mode": "either", "masters_three_years_requires_scheme": True,
            "certificate_required_levels": ["中级", "副高级", "正高级"], "textbook_core_min_words": 50000}
CONFIG = {"github_backup": {"token": "unit-test-only", "repo": cloud.PRIVATE_REPO, "branch": "main"}}


def sample_document():
    case = review.new_case("示例甲", "2026-10-09")
    return review.update_case(review.empty_cases(), case, "登记收件")


@pytest.mark.parametrize("first,second", [("符合", "不符合"), ("不符合", "符合"), ("符合", "未核验")])
def test_either_complete_scheme_suffices(first, second):
    assert review.normal_scheme_result(first, second).startswith("满足其中一套")


def test_unreviewed_or_failed_schemes_cannot_pass():
    assert review.normal_scheme_result("未核验", "未核验") == "方案仍待核验"
    assert review.normal_scheme_result("不符合", "未核验") == "方案仍待核验"
    assert review.normal_scheme_result("不符合", "不符合") == "两套方案均未满足"


@pytest.mark.parametrize("level", ["中级", "副高级", "正高级"])
def test_middle_and_senior_certificate_is_required(level):
    assert "必须有" in review.certificate_guidance(level, "无", SETTINGS)
    assert "需补证" in review.certificate_guidance(level, "无", SETTINGS)
    assert "待核验" in review.certificate_guidance(level, "待核验", SETTINGS)


def test_junior_certificate_not_required_and_50000_is_inclusive():
    assert "不强制" in review.certificate_guidance("初级", "无", SETTINGS)
    assert review.textbook_credit(49999, SETTINGS) == "发行期刊论文1篇"
    assert review.textbook_credit(50000, SETTINGS) == "核心论文1篇"
    assert review.textbook_credit(50001, SETTINGS) == "核心论文1篇"


def test_receipt_keeps_unknown_channel_and_review_distinct():
    document = sample_document()
    case = document["cases"][0]
    assert case["paper"] == case["electronic"] == "待确认"
    assert case["checks"] == {} and case["stage"] == "已收待审"
    row = review.summary_rows(document)[0]
    assert row["已核验项"] == row["待处理项"] == 0
    report = review.review_report(case, {"version": "test-1", "rules": []})
    assert "尚未逐项核验" in report and "不代替学院推荐" in report


def test_edit_preserves_other_cases_ids_and_audit():
    before = sample_document()
    before = review.update_case(before, review.new_case("示例乙", "2026-10-10"), "登记收件")
    changed = deepcopy(before["cases"][0])
    changed["level"] = "中级"
    after = review.update_case(before, changed, "更新级别")
    assert before["cases"][0]["level"] == "待确认"
    assert after["cases"][1] == before["cases"][1]
    assert after["cases"][0]["id"] == before["cases"][0]["id"]
    assert after["revision"] == before["revision"] + 1
    assert len(after["audit"]) == len(before["audit"]) + 1


def test_records_are_registered_for_guarded_private_sync():
    document = sample_document()
    assert review.CASES_PATH in MANAGED_FILES and review.RULES_PATH in MANAGED_FILES
    assert validate_backup(review.CASES_PATH, json.dumps(document)) == document
    removed = deepcopy(document)
    removed["cases"] = []
    assert would_drop_records(review.CASES_PATH, document, removed)
    duplicate = deepcopy(document)
    duplicate["cases"] += deepcopy(duplicate["cases"])
    with pytest.raises(ValueError):
        review.validate_document(duplicate, review.CASES_PATH)


def test_empty_scan_and_legacy_word_are_not_treated_as_missing():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    scan = review.extract_material("扫描件.pdf", output.getvalue())
    assert scan["metadata"]["status"] == "需人工读取"
    assert "未提取到文字" in scan["metadata"]["message"]
    legacy = review.extract_material("旧版.doc", b"legacy-data")
    assert legacy["metadata"]["status"] == "需人工读取"
    assert "不按缺件" in legacy["metadata"]["message"]


def test_partial_pdf_does_not_claim_complete_extraction(monkeypatch):
    import pypdf
    pages = [SimpleNamespace(extract_text=lambda: "有文字的第一页"), SimpleNamespace(extract_text=lambda: "")]
    monkeypatch.setattr(pypdf, "PdfReader", lambda stream: SimpleNamespace(pages=pages, is_encrypted=False))
    result = review.extract_material("部分扫描.pdf", b"test")
    assert result["metadata"]["status"] == "部分提取"
    assert "2" in result["metadata"]["message"]
    assert result["segments"][0]["location"] == "第1页"


def test_zip_reads_in_memory_and_rejects_parent_paths():
    payload = BytesIO()
    with ZipFile(payload, "w") as archive:
        archive.writestr("目录/材料.txt", "test material".encode())
    result = review.extract_upload("材料.zip", payload.getvalue())
    assert result[0]["segments"][0]["text"] == "test material"
    unsafe = BytesIO()
    with ZipFile(unsafe, "w") as archive:
        archive.writestr("../outside.txt", b"unsafe")
    with pytest.raises(ValueError):
        review.extract_upload("异常.zip", unsafe.getvalue())
    with pytest.raises(ValueError):
        review.extract_upload("损坏.zip", b"broken")


def test_docx_table_and_xlsx_comment_keep_location():
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    doc = Document()
    doc.add_paragraph("正文示例")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "需签字"
    output = BytesIO()
    doc.save(output)
    extracted = review.extract_material("表格.docx", output.getvalue())
    assert {s["location"] for s in extracted["segments"]} == {"正文段落1", "表1第1行"}
    book = Workbook()
    book.active["B2"] = "=1+2"
    book.active["B2"].comment = Comment("本人手写", "test")
    output = BytesIO()
    book.save(output)
    extracted = review.extract_material("表格.xlsx", output.getvalue())
    text = "".join(s["text"] for s in extracted["segments"])
    assert "B2==1+2" in text and "B2批注=本人手写" in text


def test_csv_names_cannot_become_spreadsheet_formulas():
    document = sample_document()
    document["cases"][0]["name"] = "=1+2"
    assert "'=1+2" in review.summary_csv(document).decode("utf-8-sig")


def test_shared_section_source_is_kept_in_individual_review():
    rule = {"id": "A03-01", "attachment": 3, "section": "附件3 报送清单", "text": "按级别核对材料。"}
    book = {"markdown": "### 附件3 报送清单\n□ 按级别核对材料。\n依据：附件3第1—5页。\n### 下一节\n"}
    assert review.rule_source(rule, book) == "附件3第1—5页。"


class StopPage(Exception):
    pass


class FakePage:
    def __init__(self):
        self.session_state = {"m30_cases": "private-data", "m30_preview": "private-material"}
        self.secrets = {}

    def info(self, *args):
        pass

    def stop(self):
        raise StopPage()

    def rerun(self):
        raise StopPage()


def page_functions(fake):
    from utils import budget_auth
    page = Path(__file__).resolve().parents[1] / "pages" / "29_30_title_review.py"
    module = ast.parse(page.read_text(encoding="utf-8"))
    functions = [node for node in module.body if isinstance(node, ast.FunctionDef)]
    namespace = {"st": fake, "AUTH_KEY": "_m30_auth", "budget_auth": budget_auth,
                 "hashlib": hashlib, "hmac": hmac, "time": time, "os": SimpleNamespace(environ={}),
                 "review": review, "private_sync": SimpleNamespace(save_document=lambda *a, **k: {"ok": False, "message": "冲突"})}
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(page), "exec"), namespace)
    # The actual page gates every top-level private read, not just its forms.
    gate_line = next(node.lineno for node in module.body if isinstance(node, ast.Expr)
                     and isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Name)
                     and node.value.func.id == "require_access")
    reads = [node.lineno for node in ast.walk(module) if isinstance(node, ast.Call)
             and isinstance(node.func, ast.Attribute) and node.func.attr == "load_document"]
    assert reads and all(line > gate_line for line in reads)
    return namespace


def test_locked_page_discards_sensitive_session_state_before_reading():
    fake = FakePage()
    namespace = page_functions(fake)
    with pytest.raises(StopPage):
        namespace["require_access"]()
    assert not any(key.startswith("m30_") for key in fake.session_state)


def test_failed_page_save_retains_a_downloadable_draft():
    fake = FakePage()
    fake.session_state["m30_cases_sha"] = "a" * 40
    namespace = page_functions(fake)
    document = sample_document()
    with pytest.raises(StopPage):
        namespace["persist"](document)
    assert fake.session_state["m30_draft"] == document
    assert fake.session_state["m30_cases"] == "private-data"


class Response:
    def __init__(self, status, payload):
        self.status_code = status
        self.payload = payload

    def json(self):
        return self.payload


class Transport:
    def __init__(self, document=None, *, private=True, conflict=False, bad_readback=False):
        self.document = document
        self.sha = "a" * 40 if document else None
        self.private = private
        self.conflict = conflict
        self.bad_readback = bad_readback
        self.puts = []
        self.file_reads = 0

    def get(self, url, **kwargs):
        if "/contents/" not in url:
            return Response(200, {"private": self.private, "full_name": cloud.PRIVATE_REPO})
        self.file_reads += 1
        if self.document is None:
            return Response(404, {})
        document = deepcopy(self.document)
        if self.puts and self.bad_readback:
            document["revision"] += 1
        return Response(200, {"sha": self.sha, "encoding": "base64",
                              "content": base64.b64encode(json.dumps(document).encode()).decode()})

    def put(self, url, *, json, **kwargs):
        self.puts.append(json)
        if self.conflict:
            return Response(409, {})
        self.document = __import__("json").loads(base64.b64decode(json["content"]))
        self.sha = "b" * 40
        return Response(200, {"content": {"sha": self.sha}})


def test_cloud_save_verifies_new_version_and_body():
    before = sample_document()
    after = review.update_case(before, before["cases"][0], "补充核验")
    transport = Transport(before)
    result = cloud.save_document(review.CASES_PATH, after, transport.sha, CONFIG, {}, transport)
    assert result["ok"] and result["sha"] == "b" * 40
    assert result["document"] == after and transport.file_reads == 2


def test_concurrent_change_never_overwrites_remote():
    before = sample_document()
    after = review.update_case(before, before["cases"][0], "补充核验")
    transport = Transport(before)
    result = cloud.save_document(review.CASES_PATH, after, "c" * 40, CONFIG, {}, transport)
    assert not result["ok"] and result["conflict"] and not transport.puts
    transport = Transport(before, conflict=True)
    result = cloud.save_document(review.CASES_PATH, after, transport.sha, CONFIG, {}, transport)
    assert not result["ok"] and result["conflict"] and len(transport.puts) == 1


def test_public_target_and_public_metadata_block_reading():
    transport = Transport(sample_document(), private=False)
    result = cloud.load_document(review.CASES_PATH, CONFIG, {}, transport)
    assert not result["ok"] and transport.file_reads == 0
    config = {"github_backup": {"token": "unit-test-only", "repo": "yaoy2/yao_1"}}
    result = cloud.load_document(review.CASES_PATH, config, {}, transport)
    assert not result["ok"] and transport.file_reads == 0


def test_bad_readback_or_removed_case_is_not_success():
    before = sample_document()
    after = review.update_case(before, before["cases"][0], "补充核验")
    transport = Transport(before, bad_readback=True)
    result = cloud.save_document(review.CASES_PATH, after, transport.sha, CONFIG, {}, transport)
    assert not result["ok"] and "未确认" in result["message"]
    removed = deepcopy(after)
    removed["cases"] = []
    transport = Transport(before)
    result = cloud.save_document(review.CASES_PATH, removed, transport.sha, CONFIG, {}, transport)
    assert not result["ok"] and not transport.puts
