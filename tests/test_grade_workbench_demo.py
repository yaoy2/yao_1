from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest

from scripts import demo_grade_workbench as demo


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "demo_grade_workbench.py"


def run_demo(output_dir, cwd):
    return subprocess.run(
        [sys.executable, "-B", str(SCRIPT), "--output-dir", str(output_dir)],
        cwd=cwd, capture_output=True, text=True, encoding="utf-8", timeout=30,
    )


def test_demo_cli_exports_verified_fictional_scores_without_business_data(tmp_path):
    output_dir = tmp_path / "demo-output"
    result = run_demo(output_dir, tmp_path)
    assert result.returncode == 0, result.stderr
    assert "Synthetic data only" in result.stdout
    assert "DEMO001=86, DEMO002=67, DEMO003=84" in result.stdout
    assert list(output_dir.iterdir()) == [output_dir / "grade-workbench-demo.xlsx"]
    assert not (tmp_path / "data").exists()

    workbook = load_workbook(output_dir / "grade-workbench-demo.xlsx", read_only=True, data_only=True)
    try:
        sheet = workbook["个人成绩计算"]
        rows = list(sheet.values)
        records = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert {row["学号"]: row["最终成绩"] for row in records} == {
            "DEMO001": 86, "DEMO002": 67, "DEMO003": 84,
        }
        assert records[1]["报告原始分"] == 90
        assert records[1]["个人贡献系数"] == 0.5
        assert records[1]["报告个人折算分"] == 45
        assert (records[0]["统一调整"], records[0]["小组调整"], records[0]["个人调整"]) == (1, 1, 2)
        assert records[2]["个人调整"] == -1
        assert workbook["小组原始评分"]["C2"].value == 80
        assert workbook["小组原始评分"]["D2"].value == 90
        assert "完全虚构" in workbook["使用说明"]["B3"].value
        assert {"规则配置", "学生与个人调整", "校验摘要", "审计日志"}.issubset(workbook.sheetnames)
    finally:
        workbook.close()


def test_demo_cli_refuses_to_overwrite_existing_output(tmp_path):
    output_dir = tmp_path / "demo-output"
    output_dir.mkdir()
    destination = output_dir / "grade-workbench-demo.xlsx"
    sentinel = b"previous user output must remain unchanged"
    destination.write_bytes(sentinel)
    result = run_demo(output_dir, tmp_path)
    assert result.returncode != 0
    assert "Refusing to overwrite existing output" in result.stderr
    assert destination.read_bytes() == sentinel
    assert list(output_dir.iterdir()) == [destination]


def test_demo_rejects_business_data_directory_before_creating_files(tmp_path, monkeypatch):
    monkeypatch.setattr(demo, "ROOT", tmp_path)
    with pytest.raises(ValueError, match="outside.*data/"):
        demo.create_demo(tmp_path / "data" / "demo-output")
    assert not (tmp_path / "data").exists()
