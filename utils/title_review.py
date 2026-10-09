"""M30 review records and bounded, read-only material text extraction.

School rules and personal records are loaded from the private data repository.
This module does not infer academic eligibility from extracted text.
"""

from copy import deepcopy
from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import PurePosixPath
import csv
import json
import re
from uuid import UUID, uuid4
from zipfile import BadZipFile, ZipFile


RULES_PATH = "data/title_review_rules.json"
CASES_PATH = "data/title_review_cases.json"
LEVELS = ("待确认", "初级", "中级", "副高级", "正高级")
TRACKS = ("待确认", "一般教师", "思政课教师", "素质教师（辅导员）")
ROUTES = ("待确认", "正常申报", "破格", "同级转评", "初聘")
RECEIPTS = ("待确认", "未收", "已收", "不适用")
STAGES = ("已收待审", "初审中", "待补材料", "待口径确认", "可提交学院复核")
CHECK_STATUSES = ("未核验", "符合", "待补材料", "待纸质核验", "待口径确认", "不适用", "存在明确矛盾")
CERTIFICATES = ("待核验", "有", "无")
SCHEME_RESULTS = ("未核验", "符合", "不符合", "不适用")
EXTRACT_STATUSES = ("已提取", "部分提取", "需人工读取", "读取失败")
MAX_DOCUMENT_BYTES = 900_000
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
LOCAL_ZONE = timezone(timedelta(hours=8))


def now():
    return datetime.now(LOCAL_ZONE).isoformat(timespec="seconds")


def _require(condition, message="职称评审记录的结构或字段不完整。"):
    if not condition:
        raise ValueError(message)


def _text(value, maximum=20000, *, required=False):
    _require(type(value) is str and len(value) <= maximum and (not required or bool(value.strip())))


def _date(value, *, optional=False):
    if optional and value == "":
        return
    _text(value, 50, required=True)
    datetime.fromisoformat(value)


def _uuid(value):
    _text(value, 36, required=True)
    _require(str(UUID(value)) == value)


def empty_cases(year=2026):
    return {"schema_version": 1, "kind": "title_review_cases", "year": year,
            "revision": 0, "cases": [], "audit": []}


def new_case(name, received_on="", year=2026):
    _text(name, 100, required=True)
    return {"id": str(uuid4()), "year": year, "name": name.strip(), "department": "",
            "level": "待确认", "track": "待确认", "route": "待确认", "stage": "已收待审",
            "received_on": received_on, "paper": "待确认", "electronic": "待确认",
            "certificate": "待核验", "normal_scheme": "未核验", "alternative_scheme": "未核验",
            "scheme_evidence": "", "note": "", "materials": [], "checks": {},
            "created_at": now(), "updated_at": now()}


def validate_document(document, repo_path):
    _require(type(document) is dict and type(document.get("schema_version")) is int and document["schema_version"] == 1)
    _require(type(document.get("year")) is int and 2000 <= document["year"] <= 2100)
    _require(len(json.dumps(document, ensure_ascii=False, allow_nan=False).encode("utf-8")) <= MAX_DOCUMENT_BYTES,
             "评审记录超过单份账本容量，未截断内容，请先导出并按年度分批整理。")
    if repo_path == RULES_PATH:
        _require(document.get("kind") == "title_review_rules")
        _text(document.get("version"), 80, required=True)
        _text(document.get("markdown"), 300000, required=True)
        _require(type(document.get("rules")) is list and 1 <= len(document["rules"]) <= 500)
        ids = []
        for rule in document["rules"]:
            _require(type(rule) is dict)
            _text(rule.get("id"), 40, required=True)
            ids.append(rule["id"])
            _text(rule.get("section"), 200, required=True)
            _text(rule.get("text"), 10000, required=True)
            _require(type(rule.get("attachment")) is int and 0 <= rule["attachment"] <= 18)
        _require(len(ids) == len(set(ids)))
        _require(type(document.get("attachments")) is list and len(document["attachments"]) == 18)
        _require(all(type(a) is dict for a in document["attachments"]))
        _require({a.get("id") for a in document["attachments"]} == set(range(1, 19)))
        for attachment in document["attachments"]:
            _text(attachment.get("filename"), 300, required=True)
            _require(bool(re.fullmatch(r"[a-f0-9]{64}", attachment.get("sha256", ""))))
        _require(type(document.get("decisions")) is list)
        decision_ids = []
        for decision in document["decisions"]:
            _require(type(decision) is dict)
            _text(decision.get("id"), 20, required=True)
            decision_ids.append(decision["id"])
            _text(decision.get("basis"), 3000)
            _text(decision.get("text"), 3000, required=True)
            _require(decision.get("status") in {"已确认", "按原文执行", "暂沿用", "个案再询问"})
        _require(len(set(decision_ids)) == len(decision_ids))
        settings = document.get("settings")
        _require(type(settings) is dict and settings.get("normal_scheme_mode") == "either")
        _require(settings.get("masters_three_years_requires_scheme") is True)
        _require(type(settings.get("textbook_core_min_words")) is int and settings["textbook_core_min_words"] > 0)
        _require(type(settings.get("certificate_required_levels")) is list
                 and set(settings["certificate_required_levels"]) <= set(LEVELS[1:]))
    elif repo_path == CASES_PATH:
        _require(document.get("kind") == "title_review_cases")
        _require(type(document.get("revision")) is int and document["revision"] >= 0)
        _require(type(document.get("cases")) is list and len(document["cases"]) <= 2000)
        ids = []
        for case in document["cases"]:
            _require(type(case) is dict)
            _uuid(case.get("id"))
            ids.append(case["id"])
            _require(case.get("year") == document["year"])
            _text(case.get("name"), 100, required=True)
            _text(case.get("department"), 200)
            _require(case.get("level") in LEVELS and case.get("track") in TRACKS and case.get("route") in ROUTES)
            _require(case.get("paper") in RECEIPTS and case.get("electronic") in RECEIPTS)
            _require(case.get("stage") in STAGES and case.get("certificate") in CERTIFICATES)
            _require(case.get("normal_scheme") in SCHEME_RESULTS and case.get("alternative_scheme") in SCHEME_RESULTS)
            _text(case.get("scheme_evidence"))
            _text(case.get("note"))
            _date(case.get("received_on"), optional=True)
            _date(case.get("created_at"))
            _date(case.get("updated_at"))
            _require(type(case.get("checks")) is dict and len(case["checks"]) <= 1000)
            for rule_id, check in case["checks"].items():
                _text(rule_id, 40, required=True)
                _require(type(check) is dict and check.get("status") in CHECK_STATUSES)
                _text(check.get("evidence"))
                _text(check.get("note"))
                _text(check.get("rule_version"), 80, required=True)
                _date(check.get("updated_at"))
            _require(type(case.get("materials")) is list and len(case["materials"]) <= 500)
            material_ids = []
            for item in case["materials"]:
                _text(item.get("id"), 64, required=True)
                material_ids.append(item["id"])
                _text(item.get("filename"), 500, required=True)
                _require(bool(re.fullmatch(r"[a-f0-9]{64}", item.get("sha256", ""))))
                _require(type(item.get("bytes")) is int and item["bytes"] >= 0)
                _require(item.get("status") in EXTRACT_STATUSES)
                _text(item.get("message"), 2000)
            _require(len(material_ids) == len(set(material_ids)))
        _require(len(ids) == len(set(ids)))
        _require(type(document.get("audit")) is list and len(document["audit"]) <= 20000)
        for entry in document["audit"]:
            _require(type(entry) is dict)
            _uuid(entry.get("case_id"))
            _date(entry.get("at"))
            _text(entry.get("action"), 300, required=True)
    else:
        raise ValueError("不支持的职称评审数据路径。")
    return document


def update_case(document, case, action):
    result = deepcopy(document)
    case = deepcopy(case)
    case["updated_at"] = now()
    existing = next((i for i, item in enumerate(result["cases"]) if item["id"] == case["id"]), None)
    if existing is None:
        result["cases"].append(case)
    else:
        result["cases"][existing] = case
    result["revision"] += 1
    result["audit"].append({"at": now(), "case_id": case["id"], "action": action})
    return validate_document(result, CASES_PATH)


def normal_scheme_result(normal, alternative):
    """Either complete scheme suffices; unknown facts can never imply a pass."""
    _require(normal in SCHEME_RESULTS and alternative in SCHEME_RESULTS)
    if "符合" in (normal, alternative):
        return "满足其中一套方案（依据人工逐项核验）"
    if normal == alternative == "不符合":
        return "两套方案均未满足"
    return "方案仍待核验"


def certificate_guidance(level, certificate, settings):
    if level == "待确认":
        return "先确认申报级别，再核对教师资格证要求。"
    if level not in settings["certificate_required_levels"]:
        return "本级别不强制要求教师资格证。"
    if certificate == "有":
        return "本级别必须有教师资格证；已登记有证，仍需核对证件。"
    if certificate == "无":
        return "本级别必须有教师资格证，当前登记为无，需补证或说明。"
    return "本级别必须有教师资格证，当前待核验。"


def textbook_credit(words, settings):
    _require(type(words) is int and words >= 0, "个人字数应为非负整数。")
    if words >= settings["textbook_core_min_words"]:
        return "核心论文1篇"
    return "发行期刊论文1篇"


def expected_materials(case):
    """Suggestions only; absence is not inferred from a filename."""
    level, route = case["level"], case["route"]
    if level == "待确认" or route == "待确认":
        return ["先确认申报级别和申报方式，再确定应交材料。"]
    rows = ["附件17个人填报并由学院汇总；身份证、学历、学位等证明按材料清单核对。"]
    if route == "同级转评":
        rows += ["附件7转评表、附件9及相应成果佐证；考核年份和高级代表作在个案中核对。"]
    elif level == "初级":
        rows += ["附件8初聘表；附件11及成果材料有则提供。"]
    else:
        rows += ["附件4申报表；附件5高级简表或附件6中级简表，按实际层级提供。",
                 "附件9高级近5学年、中级近3学年，每学年一份；附件11及成果佐证。"]
        if level in {"副高级", "正高级"}:
            rows += ["2篇代表作，每篇配一份附件10；匿名及签章要求逐项核对。"]
    if route == "破格":
        rows += ["另加附件14及破格条件证明；按专门条件审查，不套用普通冲抵。"]
    rows += ["附件12、13、15、16按实际项目、教材等成果类型使用，无相应成果不强求全交。"]
    return rows


def _zip_members(payload, max_entries=5000):
    try:
        archive = ZipFile(BytesIO(payload))
    except (BadZipFile, OSError) as error:
        raise ValueError("压缩包无法读取，请提供完整副本。") from error
    try:
        infos = archive.infolist()
        _require(len(infos) <= max_entries, "压缩包文件过多，请分批整理。")
        _require(sum(item.file_size for item in infos) <= MAX_EXPANDED_BYTES,
                 "压缩包展开体积过大，未读取，请分批整理。")
        for item in infos:
            path = PurePosixPath(item.filename.replace("\\", "/"))
            _require(not path.is_absolute() and ".." not in path.parts and not re.match(r"^[A-Za-z]:", str(path)),
                     "压缩包含异常路径，未读取。")
            _require(not item.flag_bits & 1, "压缩包已加密，请提供可读取副本。")
        return archive, infos
    except Exception:
        archive.close()
        raise


def extract_material(filename, payload):
    """No disk writes, macro execution, external requests, OCR, or silent truncation."""
    filename = str(filename).replace("\\", "/")
    _text(filename, 500, required=True)
    _require(type(payload) is bytes and len(payload) <= MAX_FILE_BYTES, "单份材料上限20 MiB，请分批提供。")
    digest = sha256(payload).hexdigest()
    metadata = {"id": sha256((filename + digest).encode()).hexdigest(), "filename": filename,
                "sha256": digest, "bytes": len(payload), "status": "需人工读取", "message": ""}
    segments = []
    partial = False
    suffix = PurePosixPath(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(payload))
            _require(not reader.is_encrypted, "PDF已加密，请提供可读取副本。")
            _require(len(reader.pages) <= 200, "PDF超过200页，请分批核阅。")
            missing = []
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    segments.append({"location": f"第{i}页", "text": text})
                else:
                    missing.append(str(i))
            partial = bool(missing)
            metadata["message"] = ("以下页面未提取到文字，需查看原件：" + "、".join(missing)) if missing else "文字供定位；签章、图片和原件需另核验。"
        elif suffix == ".docx":
            from docx import Document
            archive, _ = _zip_members(payload)
            archive.close()
            doc = Document(BytesIO(payload))
            for i, paragraph in enumerate(doc.paragraphs, 1):
                if paragraph.text.strip():
                    segments.append({"location": f"正文段落{i}", "text": paragraph.text})
            for i, table in enumerate(doc.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    text = " | ".join(cell.text for cell in row.cells)
                    if text.strip(" |"):
                        segments.append({"location": f"表{i}第{j}行", "text": text})
            metadata["message"] = "已提取正文与表格，段落号不等于页码；图片、批注及签章须对照原文件核验。"
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            archive, _ = _zip_members(payload)
            archive.close()
            workbook = load_workbook(BytesIO(payload), data_only=False, read_only=False, keep_links=False)
            try:
                _require(sum(sheet.max_row * sheet.max_column for sheet in workbook) <= 200000,
                         "工作簿范围过大，需拆分或人工核阅。")
                for sheet in workbook:
                    for row in sheet:
                        values = []
                        for cell in row:
                            if cell.value is not None:
                                values.append(f"{cell.coordinate}={cell.value}")
                            if cell.comment and cell.comment.text:
                                values.append(f"{cell.coordinate}批注={cell.comment.text}")
                        if values:
                            segments.append({"location": f"{sheet.title} 第{row[0].row}行", "text": "；".join(values)})
            finally:
                workbook.close()
            metadata["message"] = "已提取单元格及批注，保留公式原文；图片、签章和计算结果须另核验。"
        elif suffix in {".txt", ".md", ".csv"}:
            segments = [{"location": "全文", "text": payload.decode("utf-8-sig")}]
            metadata["message"] = "按UTF-8读取，仅用于材料定位。"
        else:
            metadata["message"] = "已登记文件；旧版Word/Excel、图片或其他格式需人工读取，不按缺件处理。"
        _require(sum(len(s["text"]) for s in segments) <= 1_000_000,
                 "文字量过大，未截断保存，请分批核阅。")
        if any(s["text"].strip() for s in segments):
            metadata["status"] = "部分提取" if partial else "已提取"
        elif not metadata["message"]:
            metadata["message"] = "未提取到文字，需人工核阅，不代表空白或缺件。"
    except Exception:
        metadata["status"] = "读取失败"
        metadata["message"] = "文字未成功读取，文件仍已登记；请查看原件或提供可读取副本，不按缺件处理。"
        segments = []
    return {"metadata": metadata, "segments": segments}


def extract_upload(filename, payload):
    if not str(filename).lower().endswith(".zip"):
        return [extract_material(filename, payload)]
    _require(type(payload) is bytes and len(payload) <= 50 * 1024 * 1024, "压缩包上限50 MiB，请分批提供。")
    archive, infos = _zip_members(payload, 100)
    try:
        _require(all(i.file_size <= MAX_FILE_BYTES for i in infos), "压缩包内单份文件超过20 MiB。")
        results = []
        characters = 0
        for item in infos:
            if item.is_dir():
                continue
            result = extract_material(item.filename, archive.read(item))
            characters += sum(len(s["text"]) for s in result["segments"])
            _require(characters <= 2_000_000, "本批文字量较大，请分批核阅。")
            results.append(result)
        return results
    finally:
        archive.close()


def summary_rows(document):
    rows = []
    for case in document["cases"]:
        unresolved = sum(c["status"] in {"待补材料", "待纸质核验", "待口径确认", "存在明确矛盾"}
                         for c in case["checks"].values())
        reviewed = sum(c["status"] != "未核验" for c in case["checks"].values())
        rows.append({"姓名": case["name"], "部门": case["department"], "级别": case["level"],
                     "申报方式": case["route"], "收件日期": case["received_on"], "纸质": case["paper"],
                     "电子": case["electronic"], "状态": case["stage"], "文件数": len(case["materials"]),
                     "已核验项": reviewed, "待处理项": unresolved})
    return rows


def summary_csv(document):
    rows = summary_rows(document)
    output = StringIO(newline="")
    headers = ("姓名", "部门", "级别", "申报方式", "收件日期", "纸质", "电子", "状态", "文件数", "已核验项", "待处理项")
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow({key: "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value
                         for key, value in row.items()})
    return output.getvalue().encode("utf-8-sig")


def rule_source(rule, rulebook):
    if "依据：" in rule["text"]:
        return rule["text"].split("依据：", 1)[1]
    section = rule["section"]
    active = False
    for line in rulebook.get("markdown", "").splitlines():
        if re.match(r"^#{2,3} ", line):
            if active:
                break
            active = line.lstrip("# ") == section
        elif active and line.startswith("依据："):
            return line[3:]
    if rule["attachment"]:
        return f"附件{rule['attachment']} · {section}（详见完整清单本节出处）"
    return f"完整清单 · {section}（结合本节认定条款核验）"


def review_report(case, rulebook):
    rules = {r["id"]: r for r in rulebook["rules"]}
    lines = [f"# {case['name']} 职称材料初审记录", "", f"规则版本：{rulebook['version']}",
             f"申报：{case['level']} / {case['track']} / {case['route']}",
             f"收件：{case['received_on'] or '待确认'}；纸质{case['paper']}；电子{case['electronic']}",
             f"工作状态：{case['stage']}（人工记录，不代替学院推荐或学校评审）", "",
             "## 材料清点"]
    lines += [f"- {m['filename']}：{m['status']}；{m['message']}" for m in case["materials"]] or ["尚未登记电子文件。"]
    lines += ["", "## 方案核验", normal_scheme_result(case["normal_scheme"], case["alternative_scheme"]),
              f"核验依据：{case['scheme_evidence'] or '尚未填写'}", "", "## 逐项核验"]
    for rule_id, check in case["checks"].items():
        rule = rules.get(rule_id)
        lines += [f"### {rule_id} {check['status']}", rule["text"] if rule else "历史规则，请查对应版本。",
                  "规则出处：" + (rule_source(rule, rulebook) if rule else "请查本项记录的规则版本"),
                  f"材料位置与事实：{check['evidence'] or '尚未填写'}", f"待补或复核说明：{check['note'] or '无'}",
                  f"记录规则版本：{check['rule_version']}", ""]
    if not case["checks"]:
        lines += ["尚未逐项核验；收件不等于初审通过。"]
    lines += ["", "## 备注", case["note"] or "无"]
    return "\n".join(lines) + "\n"
