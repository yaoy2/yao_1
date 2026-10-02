import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from app.config import Settings
from app.db import db_session, init_db, set_plan_approved, set_plan_execute_status, upsert_file, upsert_plan
from app.reporting import REPORT_COLUMNS, _report_row, collect_report, export_reports


class ReportingExportsTest(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmpdir.cleanup)
        self.root = Path(self.tmpdir.name)
        self.settings = Settings(
            openlist_base_url="http://127.0.0.1:1",
            openlist_username="",
            openlist_password="",
            openlist_mount_path="/云下载",
            allowed_root="/云下载",
            write_mode=False,
            default_max_files=50,
            default_scan_depth=2,
            default_scan_dir="/云下载",
            db_path=self.root / "index.sqlite",
            log_path=self.root / "test.log",
        )
        init_db(self.settings.db_path)

    def add_plan(self, number, *, approved=False, status="not_executed", **changes):
        row = {
            "file_id": str(number),
            "file_id_source": "native",
            "parent_id": "123",
            "name": f"测试 {number}.mp4",
            "full_path": f"/云下载/测试 {number}.mp4",
            "original_name": f"测试 {number}.mp4",
            "original_path": f"/云下载/测试 {number}.mp4",
            "suggested_name": f"测试 {number}.mp4",
            "suggested_path": f"普通视频/测试 {number}.mp4",
            "category": "普通视频",
            "confidence": "high",
            "size": 100,
            "reason": "按扩展名识别",
        }
        row.update(changes)
        with db_session(self.settings.db_path) as conn:
            row["file_row_id"] = upsert_file(conn, row)
            plan_id = upsert_plan(conn, row)
            set_plan_approved(conn, [plan_id], approved)
            set_plan_execute_status(conn, plan_id, status)
        return plan_id

    def test_excel_keeps_formula_and_error_like_text_literal_after_reopening(self):
        self.add_plan(
            1,
            original_name='=HYPERLINK("https://example.invalid","测试")',
            suggested_name="#REF!",
            suggested_path="=1+1",
            file_id="001234567890123456789",
            hash_sha1="#N/A",
            reason="=1+1",
        )
        self.add_plan(2, original_name="#VALUE!", suggested_name="=2+2")
        result = export_reports(self.settings, self.root / "reports")
        report = json.loads(Path(result["json"]).read_text(encoding="utf-8"))

        for data_only in (False, True):
            with self.subTest(data_only=data_only):
                workbook = load_workbook(result["xlsx"], read_only=True, data_only=data_only)
                try:
                    sheet = workbook["整理计划"]
                    for row_index, expected in enumerate(report["rows"], 2):
                        for column_index, column in enumerate(REPORT_COLUMNS, 1):
                            cell = sheet.cell(row_index, column_index)
                            value = expected[column]
                            if value != "":
                                self.assertEqual(value, cell.value, column)
                                self.assertEqual("s" if isinstance(value, str) else "n", cell.data_type)
                finally:
                    workbook.close()

    def test_report_counts_and_html_distinguish_approval_risks_and_completion(self):
        self.add_plan(1, original_name="<script>alert('测试')</script>.mp4")
        self.add_plan(2, approved=True)
        self.add_plan(3, approved=True, status="success")
        self.add_plan(4, approved=True, status="already_done")
        self.add_plan(5, approved=True, parent_id="")
        result = export_reports(self.settings, self.root / "reports")
        report = json.loads(Path(result["json"]).read_text(encoding="utf-8"))

        self.assertEqual(4, report["approved_count"])
        self.assertEqual(1, report["executable_count"])
        self.assertEqual(["否", "是", "否", "否", "否"], [row["可自动执行"] for row in report["rows"]])
        self.assertIn("缺少来源父目录ID", report["rows"][-1]["风险提示"])
        rendered = Path(result["html"]).read_text(encoding="utf-8")
        self.assertEqual(1, rendered.count('<tr class="safe">'))
        self.assertEqual(1, rendered.count('<tr class="risk">'))
        self.assertEqual(3, rendered.count('<tr class="pending">'))
        self.assertIn("&lt;script&gt;", rendered)
        self.assertNotIn("<script>", rendered)

    def test_report_blocks_invalid_names_and_parent_directories(self):
        plan_id = self.add_plan(1, approved=True)
        with db_session(self.settings.db_path) as conn:
            base = dict(conn.execute(
                "SELECT p.*, f.parent_id, f.file_id_source FROM organize_plans p "
                "JOIN files f ON f.id = p.file_row_id WHERE p.id = ?", (plan_id,)
            ).fetchone())
        cases = [
            ({"suggested_name": "../test.mp4"}, "建议文件名无效"),
            ({"suggested_name": "a" * 256}, "建议文件名无效"),
            ({"suggested_name": ""}, "建议文件名无效"),
            ({"suggested_path": "../test.mp4"}, "建议目录名无效"),
            ({"suggested_path": "a\\b/test.mp4"}, "建议目录名无效"),
            ({"suggested_path": "中" * 86 + "/test.mp4"}, "建议目录名无效"),
        ]
        for changes, reason in cases:
            with self.subTest(changes=changes):
                report_row = _report_row({**base, **changes}, {})
                self.assertEqual("否", report_row["可自动执行"])
                self.assertIn(reason, report_row["风险提示"])

    def test_failed_plan_remains_available_for_review_and_retry(self):
        self.add_plan(1, approved=True, status="failed")
        report = collect_report(self.settings)
        self.assertEqual(1, report["executable_count"])
        self.assertEqual("failed", report["rows"][0]["执行状态"])


if __name__ == "__main__":
    unittest.main()
