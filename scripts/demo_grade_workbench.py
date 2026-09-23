"""Export and verify a grading example made entirely from fictional records.

Run from the repository root:
    python scripts/demo_grade_workbench.py --output-dir outputs/demo-grade

Uses the existing pandas/openpyxl dependencies. No Streamlit, credentials,
network calls, task databases, or business data are needed.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
import shutil
import sys
from tempfile import TemporaryDirectory

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import load_workbook

from utils.grade_workbench import ScoreSettings, calculate_results, has_errors, validate_all
from utils.grade_workbench_export import build_review_workbook


OUTPUT_NAME = "grade-workbench-demo.xlsx"
EXPECTED_SCORES = {"DEMO001": 86, "DEMO002": 67, "DEMO003": 84}


def synthetic_inputs() -> tuple[pd.DataFrame, pd.DataFrame, ScoreSettings]:
    """Small examples that can be checked by hand; none represent real people."""
    students = pd.DataFrame([
        {"student_no": "DEMO001", "name": "虚构学生甲", "class_name": "虚构演示班",
         "group_code": "DEMO-A", "other_score": 80, "coefficient": 1.0,
         "individual_adjustment": 2, "adjustment_reason": "虚构示例：个人加分"},
        {"student_no": "DEMO002", "name": "虚构学生乙", "class_name": "虚构演示班",
         "group_code": "DEMO-A", "other_score": 80, "coefficient": 0.5,
         "individual_adjustment": 0, "adjustment_reason": ""},
        {"student_no": "DEMO003", "name": "虚构学生丙", "class_name": "虚构演示班",
         "group_code": "DEMO-B", "other_score": 90, "coefficient": 1.0,
         "individual_adjustment": -1, "adjustment_reason": "虚构示例：个人扣分"},
    ])
    groups = pd.DataFrame([
        {"group_code": "DEMO-A", "project_name": "虚构项目 A", "pitch_score": 80,
         "report_score": 90, "group_adjustment": 1,
         "adjustment_reason": "虚构示例：小组调整", "score_comment": "仅供演示"},
        {"group_code": "DEMO-B", "project_name": "虚构项目 B", "pitch_score": 70,
         "report_score": 80, "group_adjustment": 0,
         "adjustment_reason": "", "score_comment": "仅供演示"},
    ])
    settings = ScoreSettings(global_adjustment=1, global_adjustment_reason="虚构示例：统一调整")
    # Weights: other 60%, pitch 20%, report 20%. Apply the coefficient first.
    # DEMO001: 80*.6 + 80*.2 + 90*.2 + 1 + 1 + 2 = 86
    # DEMO002: 80*.6 + 40*.2 + 45*.2 + 1 + 1 + 0 = 67
    # DEMO003: 90*.6 + 70*.2 + 80*.2 + 1 + 0 - 1 = 84
    return students, groups, settings


def verify_workbook(path: Path) -> None:
    """Reopen the exported artifact and compare values with independent expectations."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        def records(sheet_name: str) -> list[dict]:
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = next(rows)
            return [dict(zip(headers, row)) for row in rows]

        results = records("个人成绩计算")
        scores = {row["学号"]: row["最终成绩"] for row in results}
        if len(results) != 3 or scores != EXPECTED_SCORES:
            raise RuntimeError(f"Exported scores did not match the synthetic example: {scores}")
        fields = ("路演原始分", "报告原始分", "个人贡献系数", "路演个人折算分",
                  "报告个人折算分", "统一调整", "小组调整", "个人调整", "调整合计")
        expected = {
            "DEMO001": (80, 90, 1.0, 80, 90, 1, 1, 2, 4),
            "DEMO002": (80, 90, 0.5, 40, 45, 1, 1, 0, 2),
            "DEMO003": (70, 80, 1.0, 70, 80, 1, 0, -1, 0),
        }
        for row in results:
            if tuple(row[field] for field in fields) != expected[row["学号"]]:
                raise RuntimeError(f"Exported source scores or adjustments changed: {row['学号']}")
        groups = records("小组原始评分")
        raw_scores = {
            row["小组编号"]: (row["路演原始分"], row["报告原始分"], row["小组调整"])
            for row in groups
        }
        if len(groups) != 2 or raw_scores != {"DEMO-A": (80, 90, 1), "DEMO-B": (70, 80, 0)}:
            raise RuntimeError("Exported group source scores changed.")
        if any(row["级别"] != "通过" for row in records("校验摘要")):
            raise RuntimeError("The synthetic example did not pass validation.")
        if "完全虚构" not in str(workbook["使用说明"]["B3"].value):
            raise RuntimeError("The workbook is missing its fictional-data label.")
    finally:
        workbook.close()


def create_demo(output_dir: Path) -> Path:
    output_dir = output_dir.expanduser().resolve()
    if output_dir.is_relative_to((ROOT / "data").resolve()):
        raise ValueError("Choose an output directory outside the repository's data/ folder.")
    destination = output_dir / OUTPUT_NAME
    if destination.exists() or destination.is_symlink():
        raise FileExistsError(f"Refusing to overwrite existing output: {destination}")

    students, groups, settings = synthetic_inputs()
    validation = validate_all(students, groups, settings)
    if has_errors(validation):
        raise RuntimeError("The synthetic inputs failed grading validation.")
    results = calculate_results(students, groups, settings)
    audit = pd.DataFrame([{
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "action": "生成完全虚构的演示数据",
        "detail": "DEMO001—DEMO003 均为虚构记录；没有读取实际学生或业务数据。",
    }])
    with TemporaryDirectory(prefix="yao-grade-demo-") as temporary:
        staged = Path(temporary) / OUTPUT_NAME
        build_review_workbook(
            staged,
            {"name": "完全虚构数据 / Synthetic demonstration", "term": "DEMO-TERM", "course": "DEMO-COURSE"},
            settings, students, groups, results, validation, audit,
        )
        verify_workbook(staged)
        output_dir.mkdir(parents=True, exist_ok=True)
        # Exclusive creation also protects against a file appearing after the check.
        with destination.open("xb") as target, staged.open("rb") as source:
            shutil.copyfileobj(source, target)
    verify_workbook(destination)
    return destination


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export and verify fictional grading data offline.")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "outputs" / "demo-grade",
                        help="Output directory; existing demo workbooks are never overwritten.")
    args = parser.parse_args(argv)
    try:
        destination = create_demo(args.output_dir)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"Demo failed: {exc}", file=sys.stderr)
        return 1
    print("Synthetic data only. Verified 3 fictional students after reopening the Excel workbook.")
    print("Scores: " + ", ".join(f"{student}={score}" for student, score in EXPECTED_SCORES.items()))
    print(f"Workbook: {destination}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
